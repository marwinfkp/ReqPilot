"""Text extraction for curated sources (architecture J.2): exactly txt, md, pdf, docx."""

from __future__ import annotations

from pathlib import Path

import pytest

from reqpilot.domain.errors import KnowledgeBaseError
from reqpilot.retrieval.extraction import (
    SUPPORTED_SUFFIXES,
    extract_text,
    extract_text_from_bytes,
    normalise_text,
)

pytestmark = pytest.mark.unit


def minimal_pdf(text: str) -> bytes:
    """A valid one-page PDF with one line of Helvetica text, built by hand."""
    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for offset in offsets:
        out += b"%010d 00000 n \n" % offset
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        xref,
    )
    return bytes(out)


def test_exactly_the_j2_formats_are_supported() -> None:
    assert {".txt", ".md", ".pdf", ".docx"} == SUPPORTED_SUFFIXES


def test_plain_text_and_markdown_are_read_as_utf8() -> None:
    assert extract_text_from_bytes("Clause 1 — fees".encode(), ".txt") == "Clause 1 — fees"
    assert extract_text_from_bytes(b"# Policy\n\n1. Scope", ".MD") == "# Policy\n\n1. Scope"


def test_line_endings_are_normalised_so_offsets_are_platform_independent() -> None:
    assert normalise_text("﻿a\r\nb\rc\n") == "a\nb\nc\n"
    assert extract_text_from_bytes(b"one\r\ntwo", ".txt") == "one\ntwo"


def test_pdf_text_is_extracted() -> None:
    assert "Fictional retention clause" in extract_text_from_bytes(
        minimal_pdf("Fictional retention clause"), ".pdf"
    )


def test_docx_paragraphs_are_extracted(tmp_path: Path) -> None:
    from docx import Document

    document = Document()
    document.add_paragraph("1. Fictional access clause.")
    document.add_paragraph("")
    document.add_paragraph("2. Fictional review clause.")
    path = tmp_path / "policy.docx"
    document.save(str(path))
    assert extract_text(path) == "1. Fictional access clause.\n\n2. Fictional review clause."


@pytest.mark.parametrize("suffix", [".html", ".xlsx", ".exe", ""])
def test_any_other_format_is_refused(suffix: str) -> None:
    with pytest.raises(KnowledgeBaseError, match="unsupported source format"):
        extract_text_from_bytes(b"data", suffix)


def test_a_document_with_no_text_is_refused() -> None:
    with pytest.raises(KnowledgeBaseError, match="no text"):
        extract_text_from_bytes(b"  \n ", ".txt")


def test_a_corrupt_pdf_is_refused_not_guessed_at() -> None:
    with pytest.raises(KnowledgeBaseError):
        extract_text_from_bytes(b"%PDF-1.4 this is not a pdf", ".pdf")


def test_a_missing_file_is_refused(tmp_path: Path) -> None:
    with pytest.raises(KnowledgeBaseError, match="not found"):
        extract_text(tmp_path / "absent.txt")
