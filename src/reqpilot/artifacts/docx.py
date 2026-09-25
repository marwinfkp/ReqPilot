"""DOCX rendering of a structured artefact (``FR-DOC-010``; ADR-007's python-docx path).

ADR-007 names ``pypandoc`` for Markdown -> DOCX, "with ``python-docx`` as a fallback
for environments without Pandoc". ReqPilot already depends on ``python-docx``
(architecture J.2) and adds no document framework: this renderer builds the Word
file from the **same** :class:`~reqpilot.artifacts.model.Document` the Markdown
renderer reads - same sections, blocks, citations, metadata and stamp - so the
two formats carry the same authoritative information and there is no DOCX-only
data path.

Content safety - the file is built by code, never from a template that content
could influence:

* every value is inserted as a text run, so it is **escaped by the XML writer**;
  characters XML 1.0 forbids are removed first (python-docx would refuse them);
* no field codes, no hyperlinks, no embedded objects, no images, and no macros -
  the output is a plain ``.docx`` package, never ``.docm``;
* core properties carry the title, the artefact type and the generation stamp,
  with the same stripping.
"""

from __future__ import annotations

import datetime as dt
import io
import re
import zipfile

from docx import Document as DocxDocument
from docx.shared import Pt

from reqpilot.artifacts.markdown import CONTROL_CHARS
from reqpilot.artifacts.model import Document, Fields, Items, Notice, Paragraph, Table

_SAFE_FILENAME = re.compile(r"[^A-Za-z0-9._-]+")


def clean(value: object) -> str:
    """Remove characters that cannot appear in a Word document's XML."""
    return CONTROL_CHARS.sub("", str(value))


def safe_filename(stem: str, extension: str) -> str:
    """An ASCII-only download name; project content cannot inject path or header syntax."""
    base = _SAFE_FILENAME.sub("-", stem).strip("-.")[:80] or "artefact"
    return f"{base}.{extension}"


def render_docx(document: Document, *, generated_at: dt.datetime | None = None) -> bytes:
    """Build the Word file for one document and return its bytes."""
    doc = DocxDocument()
    properties = doc.core_properties
    properties.title = clean(document.title)[:255]
    properties.subject = clean(document.artifact_type)
    properties.keywords = "ReqPilot; generated artefact"
    properties.comments = clean("; ".join(f"{label}: {value}" for label, value in document.stamp))[
        :255
    ]
    properties.author = "ReqPilot (deterministic generator)"
    stamp_time = generated_at or dt.datetime.now(dt.UTC)
    properties.created = stamp_time
    properties.modified = stamp_time
    properties.last_modified_by = properties.author
    properties.revision = 1

    doc.add_heading(clean(document.title), level=0)
    meta = doc.add_table(rows=0, cols=2)
    meta.style = "Table Grid"
    for label, value in (*document.metadata, *document.stamp):
        cells = meta.add_row().cells
        cells[0].text = clean(label)
        cells[1].text = clean(value)

    for section in document.sections:
        level = min(max(section.level - 1, 1), 4)
        doc.add_heading(f"{clean(section.number)} {clean(section.title)}", level=level)
        anchor = doc.add_paragraph()
        run = anchor.add_run(f"Section id: {clean(section.key)}")
        run.italic = True
        run.font.size = Pt(8)
        if section.citations:
            traces = doc.add_paragraph()
            traced = traces.add_run(
                "Traces to: " + ", ".join(clean(c.label) for c in section.citations)
            )
            traced.italic = True
        for block in section.blocks:
            if isinstance(block, Paragraph):
                doc.add_paragraph(clean(block.text))
            elif isinstance(block, Notice):
                doc.add_paragraph(clean(block.text), style="Intense Quote")
            elif isinstance(block, Items):
                style = "List Number" if block.ordered else "List Bullet"
                for item in block.items:
                    doc.add_paragraph(clean(item), style=style)
            elif isinstance(block, Fields):
                for field_label, value in block.pairs:
                    paragraph = doc.add_paragraph()
                    paragraph.add_run(f"{clean(field_label)}: ").bold = True
                    paragraph.add_run(clean(value))
            elif isinstance(block, Table):
                table = doc.add_table(rows=1, cols=len(block.columns))
                table.style = "Table Grid"
                for cell, column in zip(table.rows[0].cells, block.columns, strict=True):
                    cell.text = clean(column)
                for row in block.rows:
                    cells = table.add_row().cells
                    for cell, value in zip(cells, row, strict=True):
                        cell.text = clean(value)
    buffer = io.BytesIO()
    doc.save(buffer)
    return _reproducible(buffer.getvalue(), stamp_time)


def _reproducible(data: bytes, when: dt.datetime) -> bytes:
    """Re-pack the package with every entry stamped ``when``, in a fixed order.

    ``python-docx`` stamps each zip entry with the wall clock, so two exports of
    one stored version would differ byte for byte. The same version now always
    exports to the same bytes (and the same ``X-Content-SHA256``).
    """
    aware = when if when.tzinfo else when.replace(tzinfo=dt.UTC)
    stamp = max(aware.astimezone(dt.UTC).replace(tzinfo=None), dt.datetime(1980, 1, 1))
    date_time = (stamp.year, stamp.month, stamp.day, stamp.hour, stamp.minute, stamp.second)
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(data)) as source,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for name in source.namelist():
            info = zipfile.ZipInfo(name, date_time=date_time)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            target.writestr(info, source.read(name))
    return out.getvalue()


def docx_text(data: bytes) -> str:
    """Re-open a generated DOCX and return its visible text, for verification."""
    doc = DocxDocument(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)
